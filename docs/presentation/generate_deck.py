"""Generate the editable nine-slide BhuMi pitch. Requires python-pptx."""
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.oxml.xmlchemy import OxmlElement

OUT=Path(__file__).resolve().parent
r=Presentation(); r.slide_width=Inches(13.333); r.slide_height=Inches(7.5)
NAVY='112F40'; TEAL='087F73'; MINT='8DD5C2'; WHITE='FFFFFF'; BG='F5F8F8'; INK='183544'; MUTED='607781'; LINE='DCE7E5'; RED='B34D38'; PALE='FFF0E9'; GOLD='DFA66C'

def box(s,x,y,w,h,fill,line=None,radius=False):
 sh=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h))
 sh.fill.solid();sh.fill.fore_color.rgb=RGBColor.from_string(fill)
 sh._element.spPr.append(OxmlElement("a:effectLst"))
 if line: sh.line.color.rgb=RGBColor.from_string(line)
 else: sh.line.fill.background()
 if radius:
  try:sh.adjustments[0]=.12
  except Exception:pass
 return sh

def text(s,x,y,w,h,value,size=22,color=INK,bold=False,font='Arial',align=None):
 sh=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));tf=sh.text_frame
 tf.clear();tf.word_wrap=True
 tf.margin_left=tf.margin_right=0;tf.margin_top=tf.margin_bottom=0
 for i,line in enumerate(value.split('\n')):
  p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.text=line
  p.font.name=font;p.font.size=Pt(size);p.font.bold=bold;p.font.color.rgb=RGBColor.from_string(color)
  p.space_after=Pt(8)
  if align is not None:p.alignment=align
 return sh

def logo(s,x,y,dark=True,scale=1):
 s.shapes.add_picture(str(OUT.parent.parent/'public/brand/bhumi-symbol.png'),Inches(x),Inches(y),width=Inches(.45*scale),height=Inches(.45*scale))
 text(s,x+.52,y-.04,2,.5,'BhuMi',25,WHITE if dark else NAVY,True)

def base(n,kicker,title=None,dark=False):
 s=r.slides.add_slide(r.slide_layouts[6]);s.background.fill.solid();s.background.fill.fore_color.rgb=RGBColor.from_string(NAVY if dark else BG)
 text(s,.65,.32,11,.25,kicker.upper(),10,MINT if dark else TEAL,True)
 if title:text(s,.65,.89,12,1.1,title,33,WHITE if dark else NAVY,True)
 box(s,.65,7.04,12.03,.012,'365360' if dark else LINE)
 text(s,.65,7.15,9,.2,'BHUMI  /  LANDTECH 2026  /  P4 — SMART E-MUTATION & REGISTRY',8,'A5BFC8' if dark else MUTED)
 text(s,12,7.12,.65,.25,f'{n:02d} / 09',9,'A5BFC8' if dark else MUTED,align=PP_ALIGN.RIGHT)
 return s

def note(s,t):s.notes_slide.notes_text_frame.text=t

def pill(s,x,y,w,label,fill=TEAL,color=WHITE):
 box(s,x,y,w,.34,fill,radius=True);text(s,x+.1,y+.065,w-.2,.2,label,9,color,True)

def card(s,x,y,w,h,num,title,body):
 box(s,x,y,w,h,WHITE,LINE,True);text(s,x+.22,y+.22,.5,.4,num,18,TEAL,True)
 text(s,x+.22,y+.85,w-.44,.8,title,22,NAVY,True);text(s,x+.22,y+1.8,w-.44,h-1.95,body,17,MUTED)

# 1
s=base(1,'LandTech Innovation Hackathon 2026',dark=True)
logo(s,.75,.95)
pill(s,.75,2.0,3.0,'P4  /  SMART E-MUTATION & REGISTRY',fill='275564',color=MINT)
text(s,.75,2.65,8.5,1.5,'Land Ownership\nReconciliation Engine',43,WHITE,True)
text(s,.77,4.48,8.2,1,'Every land transaction should leave a traceable,\nconsistent ownership record.',23,'C5D9DF')
text(s,.77,6.2,9,.35,'Team BhuMi  ·  Daffodil International University, Bangladesh',16,MINT)
for i,(label,val) in enumerate([('DOCUMENTED INTEREST','1/2'),('TRANSFERRED INTEREST','1/4'),('REMAINING INTEREST','1/4')]):
 y=1.5+i*1.5;box(s,10.1,y,2.45,1.25,'204957',None,True);text(s,10.3,y+.18,2.05,.3,label,9,MINT,True);text(s,10.3,y+.52,2.05,.6,val,32,WHITE,True)
note(s,'20 seconds. Introduce BhuMi and the P4 category. We compare the ownership that a documented transaction supports with the mutation that follows. This is a synthetic prototype supporting human review, not an official government service. The logo is supplied by the project owner.')
# 2
s=base(2,'01 / The problem','Digital records can still disagree.')
text(s,.67,1.83,11.6,.6,'A deed and its mutation can add up — while assigning land interests to the wrong people.',22,MUTED)
card(s,.65,2.75,3.82,2.65,'01','Registered deed','Records a transfer between specific parties.')
card(s,4.76,2.75,3.82,2.65,'02','Ownership history','Shows what the transferor held and what co-owners retain.')
card(s,8.87,2.75,3.82,2.65,'03','Mutation record','May omit or misallocate the resulting shares.')
box(s,.65,5.76,12.04,.61,NAVY, None,True);text(s,.9,5.91,11.5,.4,'The hidden risk: an uninvolved co-owner loses interest even when the total stays correct.',19,WHITE,True)
ref=text(s,.67,6.55,12,.28,'Context: World Bank, Bangladesh Land Acquisition Diagnostic Review (2022). Prototype examples are fictional.',10,MUTED)
ref.click_action.hyperlink.address='https://documents1.worldbank.org/curated/en/099359509012228968/pdf/IDU05ea6dbd70b70e04c7f0b5ef06fbc25ad4670.pdf'
note(s,'35 seconds. Explain the separation between transaction records and subsequent ownership updates. The World Bank diagnostic provides institutional context for linking registration and land-record updates; it does not validate our prototype. Our specific demonstration is entirely fictional. The key question is not only whether the total area is correct, but whether each person retains the interest supported by the transaction.')
# 3
s=base(3,'02 / Our solution','Reconcile ownership, not just documents.')
text(s,.67,1.95,6.1,1.2,'BhuMi independently reconstructs the expected ownership state, then checks each party against mutation.',25,INK)
for i,(title,body) in enumerate([('Exact','Fractional shares, not rounded decisions.'),('Explainable','Named parties, exact differences and source IDs.'),('Reviewable','An officer records the next step and the reason.')]):
 y=3.4+i*.87;text(s,.75,y,1.8,.4,title,21,TEAL,True);text(s,2.6,y+.03,4.3,.6,body,17,MUTED)
box(s,7.5,2.0,5.15,4.5,NAVY,None,True)
text(s,7.85,2.4,4.4,.35,'WHAT MAKES BHUMI DIFFERENT',11,MINT,True)
text(s,7.85,3.05,4.25,1.75,'Correct total area\n≠\nCorrect ownership allocation',28,WHITE,True)
text(s,7.85,5.25,4.1,.85,'A MATCH means the supplied records reconcile. It does not establish legal title.',17,'BDD5DC')
note(s,'30 seconds. Lead with the distinction: this is an ownership reconciliation system, not a generic fraud-score dashboard. Exact arithmetic and event replay provide an independent expected state. Findings explain who differs, by how much and according to which records. Reviewers make decisions; BhuMi does not approve mutations or declare legal ownership.')
# 4
s=base(4,'03 / How it works','From documented history to a reviewable result.')
steps=[('01','Ownership history','Baseline + sourced events'),('02','Expected shares','Replay exact allocations'),('03','Mutation comparison','Check every owner'),('04','Evidence & review','Explain, review, retain')]
for i,(num,title,body) in enumerate(steps):
 x=.65+i*3.08;box(s,x,2.1,2.79,2.25,WHITE,LINE,True);pill(s,x+.2,2.32,.48,num);text(s,x+.2,2.99,2.35,.8,title,22,NAVY,True);text(s,x+.2,3.85,2.35,.4,body,14,MUTED)
 if i<3:text(s,x+2.84,2.98,.24,.5,'›',26,TEAL,True)
box(s,.65,4.85,7.35,1.58,'E7F2EE',None,True);text(s,.9,5.1,6.85,.4,'Evidence comes before certainty.',22,TEAL,True);text(s,.9,5.65,6.85,.5,'Missing sources or unclear event order prevent a false MATCH.',17,INK)
text(s,8.5,5.0,4.12,.55,'Keep three records separate',20,NAVY,True);text(s,8.5,5.62,4.12,.8,'Baseline  ·  Events  ·  Mutation\nA mutation never overwrites the expected history.',16,MUTED)
note(s,'35 seconds. Start from a sourced baseline and apply documented events in effective-date and explicit sequence order. Check that no sender transfers more than their recorded interest. Compare the reconstructed state to a separate mutation snapshot. Missing evidence or conflicting order stops a confident match. Corrections create new snapshots and runs; old calculations remain available.')
# 5
s=base(5,'04 / Live example — concept-note scenario','100 decimals. A hidden 10-decimal shift.')
text(s,.67,1.85,11.8,.62,'Fictional Rahim holds 60; Karim holds 40. Rahim sells 25 to fictional buyer Nabila.',21,MUTED)
xs=[.65,4.15,7.1,10.05];ws=[3.5,2.95,2.95,2.64]
headers=['PARTY','EXPECTED','MUTATION','DIFFERENCE']
for x,w,h in zip(xs,ws,headers):box(s,x,2.74,w,.53,NAVY);text(s,x+.2,2.9,w-.4,.22,h,11,WHITE,True)
rows=[('Rahim','35','35','0'),('Karim','40','30','−10'),('Nabila','25','35','+10'),('TOTAL','100','100','0')]
for i,row in enumerate(rows):
 y=3.27+i*.62;fill='E6F0ED' if i==3 else PALE if i in [1,2] else WHITE
 for j,(x,w,v) in enumerate(zip(xs,ws,row)):
  box(s,x,y,w,.6,fill);text(s,x+.2,y+.14,w-.4,.4,v,22,RED if j==3 and i in [1,2] else INK,i==3 or j==0)
box(s,.65,6.02,12.04,.56,TEAL,None,True);text(s,.88,6.14,11.6,.35,'Both totals balance. Karim loses 10 without participating in the sale.',22,WHITE,True)
text(s,.67,6.72,12,.22,'Run DEMO-P4-011  ·  MUTATION_WRONG_SHARE  ·  All figures are decimals equivalent of undivided interests.',10,MUTED)
note(s,'60 seconds. Open DEMO-P4-011 in the running application for this 100-decimal example. The prompt also includes DEMO-P4-001, a separate 120-decimal example with a 15-decimal discrepancy; do not mix them. In this slide, Rahim transfers 25, so expected interests are 35, 40 and 25. The mutation reports 35, 30 and 35. Both totals equal 100, but Karim is short 10 and Nabila has an unsupported extra 10. Run the actual engine and point to the owner differences and source references. A revised snapshot with shares A=7/20, B=2/5 and D=1/4 should match while the prior conflicting run remains.')
# 6
s=base(6,'05 / Key features','Built around evidence an officer can inspect.')
features=[('Exact share accounting','BigInt fractions keep thirds, sixths and transfers exact.'),('Nine focused rules','Over-transfer, missing mutation, wrong shares and broken chains.'),('Source-linked evidence','Expected versus actual shares with event and document references.'),('Review & audit trail','Request evidence, resolve or escalate with a recorded reason.')]
for i,(title,body) in enumerate(features):
 x=.65+(i%2)*6.15;y=2.06+(i//2)*1.85
 box(s,x,y,5.89,1.57,WHITE,LINE,True);box(s,x,y,.065,1.57,TEAL);text(s,x+.28,y+.23,5.25,.45,title,23,NAVY,True);text(s,x+.28,y+.85,5.25,.58,body,18,MUTED)
text(s,.72,6.25,11.9,.45,'Corrected inputs create a new result. Earlier findings and review actions remain visible.',20,TEAL,True)
note(s,'30 seconds. Summarize implemented features, not planned integrations. The nine rule categories cover exact accounting, mutation linkage and chain consistency. Explain that a revised snapshot generates a new run. Resolving a case changes workflow status, not the original calculation. Show a report or review queue during a live demo if time allows.')
# 7
s=base(7,'06 / Technology & validation','A transparent engine. A tested prototype.')
stack=[('INTERFACE','Next.js · React · Tailwind'),('DOMAIN','TypeScript · Zod · exact BigInt fractions'),('STORAGE','PostgreSQL via PGlite · Drizzle ORM')]
for i,(title,body) in enumerate(stack):
 y=2.05+i*1.1;box(s,.65,y,7.25,.91,WHITE,LINE,True);text(s,.88,y+.16,6.8,.21,title,10,TEAL,True);text(s,.88,y+.44,6.8,.36,body,20,NAVY,True)
box(s,8.3,2.05,4.39,3.7,NAVY,None,True)
for i,(num,lab) in enumerate([('27','unit tests passed'),('3','browser tests passed'),('11','synthetic scenarios checked')]):
 y=2.37+i*1.04;text(s,8.62,y,1.1,.68,num,36,MINT,True);text(s,9.85,y+.19,2.4,.55,lab,17,WHITE)
text(s,.7,5.88,7.1,.68,'Also verified: production build, type checks, lint and persistence across a server restart.',18,MUTED)
text(s,.7,6.61,12,.28,'Local measured results, not national accuracy claims. PGlite is a single-instance embedded PostgreSQL prototype.',11,MUTED)
note(s,'25 seconds. The exact engine is isolated from UI and database code. PGlite runs real PostgreSQL in-process with durable storage; this is not a hosted Supabase integration. State only measured results: 27 unit tests, 3 browser tests, 11 scenario expectations, and a restart-persistence check passed. The production build, lint and TypeScript checks passed. These results are synthetic validation, not proof of national-scale accuracy. Remote deployment status has not been verified.')
# 8
s=base(8,'07 / Impact & future plans','Support land offices with clearer exceptions.')
text(s,.67,1.85,11.9,.7,'Potential value: help reviewers focus on the exact party, share and source that need attention.',22,MUTED)
plans=[('NOW','Synthetic prototype','Demonstrate reproducible discrepancies and a traceable review workflow.'),('NEXT','Authorized pilot','Validate rules with land and legal experts using permitted data extracts.'),('LATER','Institutional integration','Add hosted storage, officer identities, bilingual UX and approved interfaces.')]
for i,(phase,title,body) in enumerate(plans):
 x=.65+i*4.12;box(s,x,2.85,3.8,2.95,WHITE,LINE,True);pill(s,x+.23,3.08,1.1,phase);text(s,x+.23,3.72,3.32,.8,title,24,NAVY,True);text(s,x+.23,4.65,3.32,1,body,18,MUTED)
text(s,.72,6.22,11.9,.55,'Complement existing e-mutation infrastructure. Do not replace official or judicial decisions.',20,TEAL,True)
note(s,'35 seconds. Frame benefits as potential, not measured savings or adopted policy. The next step is domain-expert review and an authorized institutional pilot, not unauthorized government access. A broader rollout needs institutional identity, hosted PostgreSQL, data protection and approved interfaces. Do not claim current government partnerships, national scalability or integrations. We aim to complement existing e-mutation services.')
# 9
s=base(9,'08 / Team & closing',dark=True)
text(s,.7,1.0,11.9,.65,'Digital records should tell the same ownership story.',34,WHITE,True)
text(s,.73,1.94,11.7,.7,'BhuMi shows where they disagree — and gives a human reviewer the evidence.',23,'C5D9DF')
roles=[('STUDENT 1 / TEAM LEAD','[Name · DIU ID]','Engine, data model and tests'),('STUDENT 2 / DEVELOPER','[Name · DIU ID]','Interface, API and demonstration')]
for i,(role,name,detail) in enumerate(roles):
 x=.72+i*6.1;box(s,x,3.05,5.78,1.72,'204957',None,True);text(s,x+.24,3.28,5.2,.3,role,11,MINT,True);text(s,x+.24,3.75,5.2,.45,name,24,WHITE,True);text(s,x+.24,4.3,5.2,.3,detail,15,'BAD0D8')
text(s,.77,5.11,11.8,.6,'Faculty mentor: [Name · Department]     Industry mentor: [Name · Organisation]',17,'C5D9DF')
link=text(s,.77,5.92,11.8,.45,'github.com/void-pronway/BhuMi-Land-Reconciliation',19,MINT,True)
link.click_action.hyperlink.address='https://github.com/void-pronway/BhuMi-Land-Reconciliation'
text(s,.77,6.57,11.8,.25,'AI assistance: OpenAI ChatGPT (GPT-6) and Codex. Student review and ownership remain essential.',11,'A5BFC8')
note(s,'30 seconds. Replace all placeholders with confirmed team and mentor details before submission. Invite the judges to inspect the repository and try the synthetic demo. Disclose ChatGPT GPT-6 and OpenAI Codex assistance; the team must understand, review and demonstrate the work. Close by asking for feedback on reconciliation rules and authorized pilot requirements. Total suggested speaking time across the nine slides is five minutes. Prepare for a separate three-minute Q&A.')
r.core_properties.title='BhuMi — Land Ownership Reconciliation Engine'
r.core_properties.subject='LandTech 2026 · P4 pitch deck · Synthetic prototype'
r.core_properties.author='Team BhuMi'
r.core_properties.keywords='BhuMi, P4, land records, reconciliation, synthetic prototype'
r.save(OUT/'BhuMi_Pitch_Deck.pptx')
(OUT/'Speaker_Notes.md').write_text('# BhuMi — five-minute speaker notes\n\n' + '\n\n'.join(f'## Slide {i+1}\n\n{slide.notes_slide.notes_text_frame.text}' for i,slide in enumerate(r.slides)))
print('Created',OUT/'BhuMi_Pitch_Deck.pptx')

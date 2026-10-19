"""Generate the editable ten-slide BhuMi pitch. Requires python-pptx."""
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.oxml.xmlchemy import OxmlElement

OUT=Path(__file__).resolve().parent
r=Presentation(); r.slide_width=Inches(13.333); r.slide_height=Inches(7.5)
NAVY='112F40'; TEAL='087F73'; MINT='8DD5C2'; WHITE='FFFFFF'; BG='F5F8F8'; INK='183544'; MUTED='607781'; LINE='DCE7E5'; RED='B34D38'; PALE='FFF0E9'; GOLD='DFA66C'

def box(s,x,y,w,h,fill,line=None,radius=False):
 sh=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h))
 sh.fill.solid();sh.fill.fore_color.rgb=RGBColor.from_string(fill)
 sh._element.spPr.append(OxmlElement("a:effectLst"))
 if line: sh.line.color.rgb=RGBColor.from_string(line)
 else: sh.line.fill.background()
 if radius:
  try:sh.adjustments[0]=.12
  except Exception:pass
 return sh

def text(s,x,y,w,h,value,size=22,color=INK,bold=False,font='Arial',align=None):
 sh=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));tf=sh.text_frame
 tf.clear();tf.word_wrap=True
 tf.margin_left=tf.margin_right=0;tf.margin_top=tf.margin_bottom=0
 for i,line in enumerate(value.split('\n')):
  p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.text=line
  p.font.name=font;p.font.size=Pt(size);p.font.bold=bold;p.font.color.rgb=RGBColor.from_string(color)
  p.space_after=Pt(8)
  if align is not None:p.alignment=align
 return sh

def logo(s,x,y,dark=True,scale=1):
 s.shapes.add_picture(str(OUT.parent.parent/'public/brand/bhumi-symbol.png'),Inches(x),Inches(y),width=Inches(.45*scale),height=Inches(.45*scale))
 text(s,x+.52,y-.04,2,.5,'BhuMi',25,WHITE if dark else NAVY,True)

def base(n,kicker,title=None,dark=False):
 s=r.slides.add_slide(r.slide_layouts[6]);s.background.fill.solid();s.background.fill.fore_color.rgb=RGBColor.from_string(NAVY if dark else BG)
 text(s,.65,.32,11,.25,kicker.upper(),10,MINT if dark else TEAL,True)
 if title:text(s,.65,.89,12,1.1,title,33,WHITE if dark else NAVY,True)
 box(s,.65,7.04,12.03,.012,'365360' if dark else LINE)
 text(s,.65,7.15,9,.2,'BHUMI  /  LANDTECH 2026  /  P4 — SMART E-MUTATION & REGISTRY',8,'A5BFC8' if dark else MUTED)
 text(s,12,7.12,.65,.25,f'{n:02d} / 10',9,'A5BFC8' if dark else MUTED,align=PP_ALIGN.RIGHT)
 return s

def note(s,t):s.notes_slide.notes_text_frame.text=t

def pill(s,x,y,w,label,fill=TEAL,color=WHITE):
 box(s,x,y,w,.34,fill,radius=True);text(s,x+.1,y+.065,w-.2,.2,label,9,color,True)

def card(s,x,y,w,h,num,title,body):
 box(s,x,y,w,h,WHITE,LINE,True);text(s,x+.22,y+.22,.5,.4,num,18,TEAL,True)
 text(s,x+.22,y+.85,w-.44,.8,title,22,NAVY,True);text(s,x+.22,y+1.8,w-.44,h-1.95,body,17,MUTED)


# 1 — 20 seconds
s=base(1,'LandTech Innovation Hackathon 2026 • P4',dark=True)
logo(s,.75,1.03)
text(s,.75,2.15,9.5,1.6,'One deed.\nA clearer ownership story.',42,WHITE,True)
text(s,.78,4.12,8.8,.95,'An explainable land-record reconciliation prototype\nfor Bangladesh — built for human review.',23,'C5D9DF')
pill(s,.78,5.54,4.55,'BILINGUAL EXPERIENCE  /  FULL-STACK PROTOTYPE',fill='275564',color=MINT)
text(s,.78,6.18,10.5,.42,'Mashfiqur  ·  Pronway  |  Daffodil International University',18,MINT)
s.shapes.add_picture(str(OUT.parent.parent/'public/brand/bhumi-logo.png'),Inches(10.1),Inches(2.15),width=Inches(2.45),height=Inches(2.45))
note(s,'20 seconds. Introduce the team and BhuMi. The rebuilt product combines a bilingual guided explanation with an authenticated reviewer workflow and PostgreSQL backend. It compares the ownership that documented transactions support against the recorded mutation. All demonstration records are fictional.')
# 2 — 25 seconds
s=base(2,'01 / The problem','A correct total can hide an incorrect owner.')
text(s,.67,1.9,11.8,.73,'Registered deeds and mutation records can disagree about who owns each share.',23,MUTED)
card(s,.65,2.95,3.82,2.65,'01','The deed','Who transferred an interest — and how much?')
card(s,4.76,2.95,3.82,2.65,'02','The history','What did that person hold before the transfer?')
card(s,8.87,2.95,3.82,2.65,'03','The mutation','Who appears to hold the resulting shares?')
box(s,.65,6.0,12.04,.6,NAVY,None,True)
text(s,.9,6.16,11.5,.35,'Checking only the total area can miss a loss imposed on an uninvolved co-owner.',19,WHITE,True)
note(s,'25 seconds. Explain the record mismatch without claiming a measured national error rate. A deed records a transfer, the prior history constrains what can be transferred, and the mutation records the resulting allocation. BhuMi tests whether those three accounts agree. A balanced total alone cannot answer that question.')
# 3 — 25 seconds
s=base(3,'02 / The rebuilt experience','Understand first. Review with evidence.')
features=[('Bangla + English','Switch the interface, findings and report language.'),('A guided public example','Learn the ownership flow without creating an account.'),('Light + dark themes','Responsive layouts for desktop and mobile.'),('A reviewer workspace','Assigned parcels, saved runs, cases and evidence reports.')]
for i,(title,body) in enumerate(features):
 x=.65+(i%2)*6.15;y=2.08+(i//2)*1.95
 box(s,x,y,5.89,1.7,WHITE,LINE,True)
 text(s,x+.25,y+.24,5.35,.45,title,23,NAVY,True)
 text(s,x+.25,y+.87,5.35,.65,body,18,MUTED)
text(s,.75,6.35,11.8,.37,'Designed to explain the discrepancy before asking someone to act on it.',20,TEAL,True)
note(s,'25 seconds. These are implemented features of the rebuilt version, not future plans. The public walkthrough does not require private workspace data. The reviewer experience adds individual authentication and saved evidence. Interface translations do not rewrite original source-record text. Light and dark preferences persist between visits.')
# 4 — 35 seconds
s=base(4,'03 / The P4 reconciliation engine','Reconstruct. Compare. Explain. Retain.')
steps=[('01','Ownership history','Baseline + sourced events'),('02','Expected shares','Replay exact allocations'),('03','Mutation comparison','Compare every party'),('04','Evidence & review','Record the next action')]
for i,(num,title,body) in enumerate(steps):
 x=.65+i*3.08;box(s,x,2.15,2.79,2.3,WHITE,LINE,True)
 pill(s,x+.2,2.4,.48,num);text(s,x+.2,3.02,2.35,.85,title,22,NAVY,True);text(s,x+.2,3.95,2.35,.3,body,14,MUTED)
 if i<3:text(s,x+2.84,3.05,.24,.45,'›',26,TEAL,True)
box(s,.65,4.95,7.35,1.55,'E7F2EE',None,True)
text(s,.9,5.18,6.8,.4,'Exact fractions. Explicit uncertainty.',22,TEAL,True)
text(s,.9,5.76,6.8,.5,'Missing evidence or unclear order prevents a false MATCH.',17,INK)
text(s,8.5,5.12,4.15,.6,'A correction adds a new snapshot.',18,NAVY,True)
text(s,8.5,5.77,4.15,.65,'Earlier runs and their findings remain available.',18,MUTED)
note(s,'35 seconds. The TypeScript engine uses exact BigInt fractions. It starts from the baseline and applies sourced events in effective-date and explicit sequence order, checking transfers and evidence. It compares the expected state with a separate mutation snapshot. Missing evidence prevents false certainty. Corrected snapshots create new results, while earlier calculations remain reviewable. MATCH means the supplied records reconcile; it does not establish legal title.')
# 5 — 55 seconds
s=base(5,'04 / Live example • DEMO-P4-011','100 decimals. A hidden 10-decimal shift.')
text(s,.67,1.85,11.8,.65,'Fictional Rahim holds 60; Karim holds 40. Rahim sells 25 to fictional buyer Nabila.',21,MUTED)
xs=[.65,4.15,7.1,10.05];ws=[3.5,2.95,2.95,2.64]
for x,w,h in zip(xs,ws,['PARTY','EXPECTED','MUTATION','DIFFERENCE']):
 box(s,x,2.74,w,.53,NAVY);text(s,x+.2,2.9,w-.4,.22,h,11,WHITE,True)
for i,row in enumerate([('Rahim','35','35','0'),('Karim','40','30','−10'),('Nabila','25','35','+10'),('TOTAL','100','100','0')]):
 y=3.27+i*.62;fill='E6F0ED' if i==3 else PALE if i in [1,2] else WHITE
 for j,(x,w,v) in enumerate(zip(xs,ws,row)):
  box(s,x,y,w,.6,fill);text(s,x+.2,y+.14,w-.4,.4,v,22,RED if j==3 and i in [1,2] else INK,i==3 or j==0)
box(s,.65,6.02,12.04,.56,TEAL,None,True)
text(s,.88,6.14,11.6,.35,'Both totals balance. Karim loses 10 without participating in the sale.',22,WHITE,True)
text(s,.67,6.72,12,.22,'Synthetic example • Decimals equivalent of undivided interests; not surveyed parcel boundaries.',10,MUTED)
note(s,'55 seconds including a short live interaction. Choose the 100-decimal public example or reviewer parcel DEMO-P4-011. Do not mix it with the separate 120-decimal example. After Rahim sells 25, the expected interests are 35, 40 and 25. Mutation instead shows 35, 30 and 35. Both totals equal 100, but Karim is short 10 and Nabila has an unsupported extra 10. Show the exact differences and source references. A corrected snapshot with A=7/20, B=2/5 and D=1/4 should match, while the old conflicting run remains. The shared demo may already contain changes from another visitor; rehearse first or use the deterministic public example.')
# 6 — 30 seconds
s=base(6,'05 / Backend controls','Sensitive workflows need explicit boundaries.')
features=[('Individual access','Salted password hashes, expiring sessions and parcel grants.'),('Strict server validation','Bounded inputs, exact fractions, dates and origin checks.'),('Conflict protection','Stale edits and simultaneous snapshots cannot silently overwrite history.'),('Traceable review','Named actors, recorded reasons and retained calculation results.')]
for i,(title,body) in enumerate(features):
 x=.65+(i%2)*6.15;y=2.1+(i//2)*1.9
 box(s,x,y,5.89,1.68,WHITE,LINE,True)
 text(s,x+.25,y+.22,5.3,.5,title,23,NAVY,True);text(s,x+.25,y+.86,5.3,.68,body,17,MUTED)
text(s,.75,6.33,11.8,.5,'Public judge access is limited to fictional demo parcels. Private credentials stay private.',18,TEAL,True)
note(s,'30 seconds. Explain what the backend enforces: individual accounts, administrator-only account management, parcel-scoped access, session revocation, origin checks, rate limits and strict input validation. Snapshot and review version checks protect against stale changes. Audit history is useful evidence, not tamper-proof storage against database administrators. Shared judge credentials are deliberately public and do not identify individual judges. These controls are a prototype foundation, not a production security certification.')
# 7 — 25 seconds
s=base(7,'06 / Architecture & free hosting','One application. Persistent PostgreSQL.')
for x,w,label,body in [(.65,3.3,'BROWSER','Next.js + React\nBangla / English\nGuided + reviewer views'),(4.35,4.1,'RENDER WEB SERVICE','Next.js frontend + API\nTypeScript + Zod\nExact reconciliation engine'),(8.85,3.84,'EXTERNAL DATABASE','PostgreSQL + Drizzle\nAccounts, runs, reviews\nPooled connections + TLS')]:
 box(s,x,2.2,w,2.8,WHITE,LINE,True);text(s,x+.22,2.48,w-.44,.4,label,12,TEAL,True);text(s,x+.22,3.15,w-.44,1.5,body,21,NAVY)
text(s,4.0,3.4,.3,.4,'›',25,TEAL,True);text(s,8.5,3.4,.3,.4,'›',25,TEAL,True)
text(s,.75,5.5,11.9,.55,'Free deployment path: Render web service + a suitable Neon or Supabase free tier.',20,INK,True)
text(s,.75,6.2,11.9,.48,'No persistent Render disk required. Free plans have quotas and cold starts; no 24/7 guarantee.',17,MUTED)
note(s,'25 seconds. Frontend and backend run in one Next.js service, not two deployments. Hosted persistence uses external PostgreSQL with a maximum three-connection application pool and verified remote TLS. Local development can use embedded PGlite. Neon or Supabase are deployment options; do not imply which provider the live site uses without confirming it. Free services have quotas and cold starts. Monitoring is not an availability guarantee, and the live hosting configuration has not been independently verified from this environment.')
# 8 — 25 seconds
s=base(8,'07 / Measured validation','Test the workflow, not just the calculation.')
for x,num,label in [(.65,'31','unit tests passed'),(4.77,'9','browser journeys passed'),(8.89,'11','synthetic scenarios')]:
 box(s,x,2.13,3.8,2.15,NAVY,None,True);text(s,x+.25,2.42,3.25,.85,num,48,MINT,True);text(s,x+.25,3.57,3.25,.4,label,20,WHITE)
text(s,.75,4.75,11.8,.55,'Checked: bilingual demo, permissions, validation, review history and concurrent edits.',20,INK,True)
text(s,.75,5.52,11.8,.55,'Also verified: PostgreSQL adapter, restart persistence, production build and type checks.',19,MUTED)
text(s,.75,6.35,11.8,.35,'Synthetic local validation • Not a claim of legal accuracy, production security or national scale.',14,TEAL)
note(s,'25 seconds. As recorded in docs/validation.md: 31 unit tests and nine browser journeys passed on the embedded adapter. The eight core journeys were also checked on native PostgreSQL during the rebuild, and the added judge journey was checked separately against PostgreSQL. Restart persistence was verified on both adapters. The subsequent branding update passed the build, lint and two relevant browser journeys. Do not claim every suite was rerun after every visual edit or that remote CI and production hosting were verified. These tests are synthetic validation, not a national accuracy estimate.')
# 9 — 25 seconds
s=base(9,'08 / Impact & next steps','Make discrepancies easier to act on.')
plans=[('IMPLEMENTED','A working prototype','Bilingual UX, exact reconciliation, individual accounts and PostgreSQL support.'),('NEXT','Domain validation','Review rules and Bangla terminology with land and legal specialists.'),('PILOT GATE','Operational readiness','Authorized data, security review, backups and institutional access controls.')]
for i,(phase,title,body) in enumerate(plans):
 x=.65+i*4.12;box(s,x,2.42,3.8,3.35,WHITE,LINE,True);pill(s,x+.23,2.7,1.65,phase)
 text(s,x+.23,3.39,3.32,.8,title,24,NAVY,True);text(s,x+.23,4.4,3.32,1.15,body,18,MUTED)
text(s,.75,6.15,11.8,.6,'Potential impact: help reviewers identify the exact person, share and source needing attention.',20,TEAL,True)
note(s,'25 seconds. Distinguish what exists from what remains. Bilingual UX, accounts and hosted PostgreSQL support are implemented. Next steps include domain review, an authorized pilot and operational controls. We have not measured officer time savings or national impact. No government integration, institutional adoption or legal verification is claimed. BhuMi should complement existing e-mutation processes and human decisions.')
# 10 — 35 seconds
s=base(10,'09 / Team & judge access','Try the record. Inspect the evidence.',dark=True)
for i,(name,identifier,role) in enumerate([('Mashfiqur','242-15-532','Engine, data model and tests'),('Pronway','242-15-335','Interface, API and demonstration')]):
 x=.7+i*6.1;box(s,x,2.05,5.8,1.55,'204957',None,True)
 text(s,x+.25,2.29,5.25,.45,name,26,WHITE,True);text(s,x+.25,2.84,5.25,.28,'DIU ID '+identifier,16,MINT);text(s,x+.25,3.24,5.25,.23,role,13,'C5D9DF')
link=text(s,.77,4.0,11.8,.5,'bhumi.diu.my.id',28,MINT,True);link.click_action.hyperlink.address='https://bhumi.diu.my.id/?view=login'
text(s,.77,4.68,11.8,.65,'Click “Enter judge demo” or use the public credentials:\njudge@bhumi.demo  /  BhuMi-Judge-Demo-2026!',18,WHITE)
link=text(s,.77,5.68,11.8,.4,'github.com/void-pronway/BhuMi-Land-Reconciliation',18,MINT);link.click_action.hyperlink.address='https://github.com/void-pronway/BhuMi-Land-Reconciliation'
text(s,.77,6.23,11.8,.26,'Shared fictional records only. Do not enter personal information.',14,'C5D9DF')
text(s,.77,6.65,11.8,.2,'AI-assisted development: OpenAI ChatGPT and Codex • Team review and ownership remain essential.',10,'A5BFC8')
note(s,'35 seconds. Preserve the names and DIU IDs supplied in the uploaded deck. Invite judges to the live URL and repository. The shared account can run checks and save reviews on fictional parcels; its changes are visible to other visitors. It cannot manage users or change its password. Verify the live deployment and rehearse before presenting; this environment could not independently reach the site. If it is unavailable, use the editable numerical example on slide 5. Confirm team roles before submission. No mentor names were supplied, so none are invented. Total speaking time is approximately five minutes, with a separate three-minute Q&A. Student understanding and review of AI-assisted work remain essential.')
r.core_properties.title='BhuMi — Rebuilt Full-Stack Prototype'
r.core_properties.subject='LandTech 2026 • P4 • Updated bilingual product pitch'
r.core_properties.author='Team BhuMi'
r.save(OUT/'BhuMi_Updated_Pitch_Deck.pptx')
(OUT/'Updated_Speaker_Notes.md').write_text('# BhuMi — updated five-minute pitch\n\n'+'\n\n'.join(f'## Slide {i+1}\n\n{slide.notes_slide.notes_text_frame.text}' for i,slide in enumerate(r.slides)))
print('Created editable deck:',OUT/'BhuMi_Updated_Pitch_Deck.pptx')
